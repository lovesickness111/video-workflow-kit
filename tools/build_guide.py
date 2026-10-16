"""Generate an optional starter draft. The manually edited DOCX is authoritative."""
from pathlib import Path
from datetime import date
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
SHOTS=ROOT/'docs/screenshots'
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output',type=Path,default=ROOT/'local/guide-draft.docx')
args=parser.parse_args()
OUT=args.output.resolve()
if OUT==(ROOT/'docs/Huong-dan-su-dung.docx').resolve():
    parser.error('The hand-edited guide is authoritative. Choose a separate draft output to preserve manual content and screenshots.')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Calibri')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.13
doc.styles['Title'].font.size=Pt(28)
doc.styles['Heading 1'].font.size=Pt(20)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(9)
for h in ['Heading 1','Heading 2']:
    doc.styles[h].paragraph_format.space_before=Pt(8)
    doc.styles[h].paragraph_format.space_after=Pt(6)
doc.core_properties.title='Hướng dẫn chỉnh sửa và đăng video bằng Codex'
doc.core_properties.subject='Thiết lập Video Workflow Kit trên Windows'
doc.core_properties.author='Video Workflow Kit'
doc.core_properties.keywords='Codex video-use ElevenLabs Zernio Windows'

def p(text='',bold=False):
    par=doc.add_paragraph()
    run=par.add_run(text);run.bold=bold
    return par

def h(text,level=2): doc.add_heading(text,level)
def page(title): doc.add_page_break();h(title,1)
def step(n,title,text):
    p(f'{n}. {title}',True);p(text)
def code(text):
    for line in text.splitlines():
        par=doc.add_paragraph()
        par.paragraph_format.space_after=Pt(3)
        run=par.add_run(line);run.font.name='Consolas';run.font.size=Pt(9)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=txt
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trPr.append(no)
        if i==0: trPr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j])
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF5');tcPr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(3);par.paragraph_format.line_spacing=1.05
                for run in par.runs:run.font.size=Pt(10);run.bold=i==0
    p()

def screenshot(name,label,detail):
    file=SHOTS/name
    if not file.exists():
        # A clearly marked illustration placeholder, never a fabricated service screenshot.
        file=ROOT/'local/doc-assets'/name
        file.parent.mkdir(parents=True,exist_ok=True)
        im=Image.new('RGB',(1600,410),'#F0F2F4')
        draw=ImageDraw.Draw(im)
        font_path=Path('C:/Windows/Fonts/arial.ttf')
        font=ImageFont.truetype(str(font_path),44) if font_path.exists() else ImageFont.load_default()
        small=ImageFont.truetype(str(font_path),30) if font_path.exists() else font
        draw.text((50,85),'VỊ TRÍ BỔ SUNG ẢNH CHỤP MÀN HÌNH',fill='#202830',font=font)
        draw.text((50,170),label,fill='#202830',font=small)
        draw.text((50,240),name,fill='#52606D',font=small)
        im.save(file)
    with Image.open(file) as im:
        w,hgt=im.size
    width=min(6.65,1.8*w/hgt)
    par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=par.add_run();run.add_picture(str(file),width=Inches(width))
    inline=run._r.xpath('.//wp:docPr')
    if inline:inline[0].set('descr',label+' — '+detail)
    cap=doc.add_paragraph(label+' — '+detail,style='Caption')
    cap.paragraph_format.keep_with_next=False
    par.paragraph_format.keep_with_next=True

doc.add_heading('Hướng dẫn chỉnh sửa và đăng video bằng Codex',0)
p('Video Workflow Kit cho người dùng Windows',True)
p('Phiên bản 0.1.0 • Ngày 16 tháng 10 năm 2026')
p('Tài liệu này giúp bạn thiết lập một lần, chỉnh video đã quay và đăng bản đã duyệt lên các tài khoản mạng xã hội của mình. Bạn làm việc bằng cách trò chuyện với Codex; các lệnh trong tài liệu có thể nhờ Codex thực hiện, ngoại trừ bước tự đăng nhập và nhập API key.')
h('Bạn sẽ làm theo thứ tự này')
p('Mở workspace → cài bộ công cụ → tạo và lưu API key → kết nối kênh → chỉnh video → xem và duyệt → đăng bài → kiểm tra kết quả.')
h('Cần chuẩn bị')
table(['Chuẩn bị','Mục đích'],[
('Máy Windows và Internet','Có quyền cài Python 3.11+, Node.js 22+, Git và FFmpeg; setup hỗ trợ cài qua winget.'),
('Codex đã đăng nhập','Dùng tài khoản có quyền sử dụng Codex. Không bắt buộc có OpenAI API key riêng khi đăng nhập bằng ChatGPT.'),
('Tài khoản ElevenLabs','Dùng Speech to Text để tạo transcript có thời gian từng từ; cần quota phù hợp.'),
('Tài khoản Zernio và kênh của bạn','Cho phép bộ công cụ đăng lên các kênh đã kết nối. Số key tùy cấu hình; không cố định hai key.'),
('Video và nhạc được phép sử dụng','Đặt trong folder dự án; nhạc nền là tùy chọn, cần quyền sử dụng phù hợp.')
],[5,12])
p('Bạn tự duyệt quyền truy cập tài khoản và nội dung trước khi đăng. Setup không tự phiên âm hoặc đăng thử. Giá và quota dịch vụ có thể thay đổi; xem trên tài khoản của bạn trước khi dùng.')

page('1 Mở workspace và cài bộ công cụ')
step(1,'Nhận bộ kit','Nếu dùng GitHub private, đăng nhập tài khoản đã được cấp quyền rồi tải ZIP phát hành. Giải nén vào một folder cố định, ví dụ C:\\VideoWorkflow. Không chạy trực tiếp bên trong file ZIP.')
step(2,'Mở folder trong Codex','Chọn mở project/folder và trỏ vào folder chứa START-HERE.md, AGENTS.md và skills.lock.json. Giữ cửa sổ chat làm việc trong project này để Codex nhận các skill của bộ kit.')
screenshot('codex-workspace.png','Ảnh 01 Codex mở workspace','Chụp folder bộ kit và nơi nhập prompt setup.')
step(3,'Gửi yêu cầu thiết lập','Dán câu dưới vào Codex. Cài phần mềm có thể mở hộp thoại của Windows; hoàn tất hộp thoại trên máy nếu xuất hiện.')
p('Đọc START-HERE.md và thiết lập bộ workflow này trên Windows cho tôi. Chưa transcribe hoặc đăng video.',True)
p('Lệnh tương ứng nếu cần chạy trong terminal tại folder bộ kit:')
code('powershell -ExecutionPolicy Bypass -File .\\setup\\setup.ps1')
p('Kết quả mong đợi: có môi trường .venv, công cụ publisher đã build và báo cáo trạng thái. Thiếu API key hoặc chưa chọn tài khoản ở giai đoạn này là bình thường. Nếu công ty chặn cài phần mềm, nhờ IT cài các thành phần rồi chạy lại setup.')

page('2 Tạo API key ElevenLabs')
step(1,'Mở danh sách API key','Đăng nhập ElevenLabs, vào Developers rồi API Keys. Có thể mở https://elevenlabs.io/app/developers/api-keys. Nếu giao diện đổi, tìm mục API Keys trong phần dành cho nhà phát triển.')
screenshot('elevenlabs-api-keys.png','Ảnh 02 Danh sách API key ElevenLabs','Chụp đường dẫn menu và nút tạo key; không lộ key.')
step(2,'Tạo key dành riêng cho phiên âm','Chọn Create API Key, đặt tên dễ nhớ như video-workflow-stt. Bật Restrict Key và cho phép Speech to Text. Với workflow phiên âm hiện tại, không cần bật Text to Speech, Sound Effects hoặc Music Generation. Nếu sau này muốn tạo giọng đọc hay nhạc bằng ElevenLabs, cần quyền tương ứng.')
screenshot('elevenlabs-stt-permission.png','Ảnh 03 Quyền Speech to Text','Chụp form tạo key với Speech to Text ở Access và các quyền không dùng ở No Access.')
step(3,'Lưu key vào bộ kit','Tạo key, sao chép và nhập vào ô ELEVENLABS_API_KEY bằng configure.ps1 theo mục 4. Không dán key vào chat hoặc ảnh hướng dẫn. Kiểm tra quota trên ElevenLabs trước lần phiên âm đầu tiên.')
p('Không có key vẫn có thể cắt, ghép hoặc trộn âm thanh bằng FFmpeg. Video-use cần transcript có timestamp từng từ để edit theo lời thoại; key STT phục vụ bước tạo transcript này. Transcript đã có có thể được tái sử dụng.')

page('3 Kết nối tài khoản mạng xã hội trên Zernio')
step(1,'Đăng nhập Zernio','Mở https://zernio.com và đăng nhập tài khoản của bạn. Nếu cần, tạo profile cho thương hiệu/nhóm theo giao diện Zernio. Chỉ kết nối những kênh bạn có quyền quản lý.')
step(2,'Kết nối nền tảng muốn dùng','Trong khu vực kết nối tài khoản, chọn YouTube, TikTok, Facebook hoặc LinkedIn. Tự đăng nhập và duyệt màn hình cấp quyền. Với Facebook, kiểm tra đúng Page; với YouTube, kiểm tra đúng kênh; với LinkedIn, phân biệt trang cá nhân và tổ chức.')
screenshot('zernio-connect-accounts.png','Ảnh 04 Chọn nền tảng trên Zernio','Chụp nơi bắt đầu kết nối một tài khoản mạng xã hội.')
step(3,'Kiểm tra tài khoản đã kết nối','Sau khi cấp quyền, quay lại Zernio và xác nhận tên kênh/Page. Bộ kit cho phép đăng một hoặc nhiều nền tảng; không yêu cầu kết nối đủ bốn. Giới hạn số tài khoản phụ thuộc gói dịch vụ của bạn.')
screenshot('zernio-connected-accounts.png','Ảnh 05 Các tài khoản đã kết nối','Chụp đủ tên nền tảng và tên kênh/Page để phân biệt đích đăng.')
p('Đăng nhập Chrome sẵn không thay thế bước kết nối Zernio. Bộ kit đăng bài qua API của Zernio; không cần sao chép Chrome profile hay cookie từ máy người khác.')

page('4 Tạo và lưu key Zernio')
step(1,'Tạo API key','Trong Zernio, mở API Keys và chọn Create API key. Đặt tên riêng cho bộ workflow. Nếu có lựa chọn giới hạn phạm vi, bảo đảm key được đọc tài khoản/health, tải media, xác thực bài đăng, tạo/đọc bài và retry bài lỗi theo nhu cầu; không cần cấp quyền quản trị không liên quan.')
screenshot('zernio-create-key.png','Ảnh 06 Tạo key Zernio','Chụp trang API Keys và nút tạo; che toàn bộ giá trị key nếu xuất hiện.')
step(2,'Mở terminal trong folder bộ kit','Yêu cầu Codex mở terminal hoặc dùng terminal của bạn. Tự chạy lệnh dưới để nhập key ẩn. Nhấn Enter ở một mục sẽ giữ giá trị cũ hoặc bỏ qua nếu chưa có.')
code('powershell -ExecutionPolicy Bypass -File .\\setup\\configure.ps1')
step(3,'Nhập hai loại key','Dán key ElevenLabs vào ELEVENLABS_API_KEY và key Zernio vào ZERNIO_API_KEY. Khi nhập, terminal không hiển thị ký tự. Đây là hành vi bình thường. Mỗi tài khoản dịch vụ phải dùng đúng key của dịch vụ đó.')
screenshot('local-key-input.png','Ảnh 07 Nhập key cục bộ','Chụp terminal ở ô nhập ẩn. Không để key plaintext trong ảnh hoặc lịch sử lệnh.')
p('Key được lưu trong local/secrets.env trên máy của bạn. File này không được đưa vào Git hoặc ZIP phát hành, nhưng vẫn là dữ liệu plaintext trên ổ đĩa. Giữ folder riêng tư. Zernio có thể chỉ hiển thị key lúc tạo; lưu bằng cách nhập cục bộ ngay sau khi tạo.')

page('5 Chọn đúng tài khoản đăng')
p('Có key mới chỉ cho phép kết nối dịch vụ. Bạn cần xác nhận đích đăng để tránh nhầm kênh hoặc Page. Thực hiện một lần và cập nhật lại khi thay tài khoản.')
step(1,'Nhờ Codex liệt kê tài khoản','Gửi: “Liệt kê tài khoản đã kết nối trên Zernio bằng công cụ publisher. Chưa đăng bài.” Lệnh này chỉ đọc tên và account ID, không đăng video.')
code('node tools/publisher/dist/cli.js accounts')
step(2,'Chọn tài khoản của bạn','Nói rõ kênh/Page muốn dùng. Codex lưu account ID vào local/accounts.json. Nếu nhiều tài khoản cùng nền tảng, chọn bằng ID cụ thể. Không sao chép file tài khoản từ máy tác giả bộ kit.')
p('Ví dụ cấu hình chỉ YouTube dưới đây dùng ID minh họa; Codex phải thay bằng ID thật vừa đọc:')
code('{"version":1,"profiles":{"main":{\n  "apiKeyEnv":"ZERNIO_API_KEY",\n  "platforms":{"youtube":{"accountId":"ID_KENH_CUA_BAN"}}\n}}}')
step(3,'Kiểm tra cấu hình','Nhờ Codex chạy check để xác minh các tài khoản đã chọn còn đăng được. Nếu cần nhiều key Zernio, nhờ người bảo trì thêm nhóm cấu hình; mỗi nền tảng đã chọn chỉ thuộc một nhóm key trong bản này.')
code('node tools/publisher/dist/cli.js check\npowershell -ExecutionPolicy Bypass -File .\\setup\\doctor.ps1')
h('Đọc kết quả kiểm tra')
p('doctor kiểm tra phần mềm, dependencies, bộ lọc phụ đề, sự có mặt của key và cấu hình. Nó không gọi ElevenLabs và không chứng minh key STT hợp lệ. check gọi Zernio để xác minh tài khoản; lần phiên âm thật mới sử dụng quota STT. Không tự thử bằng cách đăng bài công khai.')

page('6 Chỉnh video và thêm nhạc')
step(1,'Đưa video vào dự án','Tạo projects/ten-du-an/input/ bên trong bộ kit và đặt video gốc ở đó. Không đặt video trong folder skill. Giữ bản gốc; kết quả edit sẽ ở projects/ten-du-an/edit/.')
step(2,'Yêu cầu Codex đề xuất phương án','Ví dụ: “Dùng video-workflow để xem video trong projects/demo/input. Đề xuất cắt đoạn lặp, bỏ khoảng im lặng, giữ giọng nói tự nhiên, thêm phụ đề tiếng Việt. Chưa render cho tới khi tôi duyệt.”')
step(3,'Duyệt phương án','Xem độ dài mục tiêu, các đoạn giữ/bỏ, tỷ lệ khung hình và phong cách phụ đề. Chỉ yêu cầu phiên âm khi bạn đồng ý gửi phần âm thanh đến ElevenLabs và sử dụng quota. Codex tái sử dụng transcript đã có để tránh phiên âm lại không cần thiết.')
step(4,'Yêu cầu nhạc nếu cần','Ví dụ: “Thêm nhạc nền lo-fi chill nhỏ dưới giọng nói. Cho tôi nghe lựa chọn trước khi chọn nhạc.” Có thể cung cấp file nhạc của bạn. Nếu tải nhạc từ nguồn khác, lưu tên track, nguồn và giấy phép trong dự án.')
step(5,'Xem bản dựng','Mở MP4 Codex trả về. Xem đầu/cuối, điểm cắt, phụ đề tiếng Việt và nghe cân bằng nhạc với lời thoại. Nếu sai, mô tả vị trí cần sửa hoặc mốc thời gian; duyệt bản cuối trước khi yêu cầu đăng.')
h('Tiếp tục vào ngày khác')
p('Mở lại cùng workspace và nói: “Tiếp tục dự án demo từ project.md, dùng transcript và bản dựng hiện có.” Không tạo project mới để tiếp tục một lần đăng đang dở.')
h('Các file nên giữ')
p('Video gốc, transcript, EDL, phụ đề, project.md và bản MP4 cuối giúp tiếp tục edit. Lịch sử đăng ở local/runs giúp tiếp tục publish. Khi chuyển sang máy khác, sao lưu dự án và lịch sử theo kênh riêng; ZIP source không chứa chúng. Nhập key lại trên máy mới.')

page('7 Duyệt và đăng video')
step(1,'Yêu cầu chuẩn bị bài đăng','Ví dụ: “Chuẩn bị đăng bản final.mp4 của dự án demo lên kênh YouTube đã cấu hình và Facebook Page của tôi. Viết caption để tôi duyệt. YouTube để public, Facebook đăng feed. Chưa đăng ngay.”')
step(2,'Xem bản duyệt cụ thể','Codex nhập video vào một run, chuẩn bị metadata, tải media lên Zernio và xác thực payload. Bước prepare chưa tạo bài đăng, nhưng đã gửi video đến kho lưu trữ của dịch vụ. Xem video, caption, tên và ID tài khoản, quyền riêng tư và thời điểm đăng.')
screenshot('publish-review.png','Ảnh 08 Bản duyệt trước khi đăng','Chụp bản xem trước và nội dung duyệt; không chụp URL upload hoặc API key.')
step(3,'Duyệt đăng','Khi thông tin đúng, nói: “Duyệt đăng ngay cấu hình vừa chuẩn bị.” Nếu hẹn giờ, nêu ngày, giờ và múi giờ Việt Nam. Không dùng lệnh xác nhận cho một cấu hình bạn chưa xem.')
step(4,'Kiểm tra kết quả','Codex trả kết quả riêng theo nền tảng và link bài khi API cung cấp. Published là đã đăng; scheduled là đã lên lịch; publishing là còn xử lý; partial là một phần thành công. Nếu chưa có link, kiểm tra tiếp run hiện tại thay vì tạo bài mới.')
p('Lưu ý nhãn nội dung AI: chỉ dùng Codex cắt video người thật không tự động có nghĩa toàn bộ video là nội dung tổng hợp. Chọn khai báo theo nội dung thực và quy định hiện hành của nền tảng. Xác nhận thêm đối tượng trẻ em cho YouTube và quyền riêng tư TikTok.')

page('8 Tiếp tục khi bị gián đoạn')
h('Mạng lỗi hoặc đóng Codex giữa chừng')
p('Mở lại workspace và yêu cầu: “Kiểm tra run đăng gần nhất của dự án demo. Đọc trạng thái đã lưu và kiểm tra Zernio trước khi làm tiếp. Không tạo bài mới nếu kết quả lần trước chưa rõ.”')
p('Công cụ giữ request ID, post ID và trạng thái từng nhóm tài khoản. Bài đã thành công không cần đăng lại để sửa nhóm khác. Nếu chưa biết máy chủ đã nhận bài chưa, phải đối chiếu Zernio; việc chờ một lúc không chứng minh bài chưa được tạo.')
h('Một nền tảng lỗi')
p('Đọc lỗi cụ thể: hết quota, mất kết nối, video không phù hợp hoặc hạn chế từ nền tảng. Khắc phục nguyên nhân rồi duyệt retry cho bài đã ghi nhận. Không thay caption hoặc tạo run mới chỉ để né phát hiện trùng. TikTok hết khả năng đăng trực tiếp sẽ được báo lại; bản kit này không tự chuyển sang Inbox draft.')
h('Run bị khóa')
p('Khóa ngăn hai tiến trình đăng cùng run. Nhờ Codex kiểm tra tiến trình còn chạy hay đã dừng. Sau sự cố, đối chiếu kết quả Zernio trước khi xóa khóa cũ. Không tự xóa khóa chỉ vì muốn thử lại nhanh.')
h('Giữ dữ liệu khi cập nhật')
p('Chờ edit/render/publish kết thúc. API key nằm trong local/secrets.env, dự án trong projects và lịch sử đăng trong local/runs. Cập nhật một thành phần sẽ giữ các folder này. Rollback skill không xóa bài đã đăng hoặc hủy lịch đã tạo trên mạng xã hội.')

page('9 Cập nhật riêng từng skill')
p('Người bảo trì kiểm tra thay đổi upstream và phát hành bản kit đã duyệt. Bạn chỉ cập nhật theo tag hoặc ZIP được người bảo trì thông báo. Không cần tự theo dõi code video-use hoặc chạy git pull trong folder skill.')
step(1,'Nhận phiên bản đã được duyệt','Ví dụ prompt: “Cập nhật riêng video-use từ bản kit mà người bảo trì vừa cung cấp. Giữ API key, dự án và lịch sử đăng. Báo kết quả và bản sao lưu.”')
step(2,'Cập nhật từ ZIP đã tải','Giải nén bản kit mới sang một folder khác. Thay đường dẫn trong ví dụ bằng folder mới thực tế; lệnh chạy từ bộ kit đang dùng.')
code('powershell -ExecutionPolicy Bypass -File .\\setup\\update.ps1 `\n  -Component video-use -Source "C:\\VideoWorkflowMoi"')
step(3,'Hoặc cập nhật bằng tag GitHub','Chỉ dùng tag người bảo trì đã duyệt. Máy cần Git và quyền đọc repo private. v1.1.0 dưới đây là ví dụ, không phải một bản đã phát hành.')
code('powershell -ExecutionPolicy Bypass -File .\\setup\\update.ps1 `\n  -Component video-use -Version v1.1.0')
step(4,'Kiểm tra và quay lại nếu cần','Nhờ Codex đọc báo cáo cập nhật, chạy doctor và mở lại Codex nếu skill mới chưa xuất hiện. Để quay về bản đã sao lưu:')
code('powershell -ExecutionPolicy Bypass -File .\\setup\\update.ps1 `\n  -Component video-use -Rollback')
p('Nếu muốn cập nhật phần đăng bài, dùng Component publisher; skill điều phối dùng workflow. Thay đổi script setup/updater hoặc cấu trúc bộ kit cần hướng dẫn nâng cấp toàn bộ từ người bảo trì. Updater sẽ dừng nếu phát hiện bạn đã sửa trực tiếp file của thành phần.')

page('10 Xử lý lỗi và nơi cần xem')
table(['Hiện tượng','Việc cần làm'],[
('Skill không xuất hiện','Mở đúng folder chứa START-HERE.md. Khởi động lại Codex nếu vừa cài/cập nhật. Kiểm tra có bản skill trùng ở phạm vi người dùng.'),
('Không tìm thấy Python hoặc FFmpeg','Chạy lại setup; mở terminal mới sau cài. Nếu công ty chặn cài, nhờ IT hỗ trợ. doctor phải thấy ffmpeg, ffprobe và subtitles.'),
('ElevenLabs từ chối','Kiểm tra key còn hiệu lực, quyền Speech to Text và quota. Nhập lại qua configure; không gửi key vào chat hỗ trợ.'),
('Zernio không thấy kênh','Kết nối đúng tài khoản/Page trên Zernio; chạy accounts, cập nhật ID và chạy check.'),
('Đăng lỗi hoặc chưa có link','Giữ run cũ. Chạy inspect, đọc lỗi từng nền tảng; không đăng lại toàn bộ.'),
('Updater báo file đã sửa','Giữ bản chỉnh sửa; nhờ người bảo trì đối chiếu. Không xóa file để vượt kiểm tra.'),
('Chuyển sang máy khác','Tải bộ kit, setup lại, nhập key cục bộ và chọn đúng tài khoản. Sao lưu dự án và lịch sử đăng riêng nếu cần tiếp tục.')
],[5.2,11.8])
h('Tài liệu và ảnh cần bổ sung')
p('Trong Word, nhấp phải ảnh placeholder rồi chọn Change Picture để thay ảnh bạn chụp. Hoặc đặt PNG đúng tên trong docs/screenshots và nhờ người bảo trì tạo lại tài liệu. Danh sách ảnh nằm trong docs/screenshots/README.md. Che toàn bộ key và mã đăng nhập trước khi thêm ảnh.')
h('Nguồn đối chiếu')
p('Codex skills: https://learn.chatgpt.com/docs/build-skills\nCodex đăng nhập: https://learn.chatgpt.com/docs/auth\nVideo-use: https://github.com/browser-use/video-use\nElevenLabs API keys: https://elevenlabs.io/docs/overview/administration/workspaces/api-keys\nZernio bắt đầu: https://docs.zernio.com/\nZernio media: https://docs.zernio.com/guides/media-uploads')
p('Menu dịch vụ có thể thay đổi theo thời gian. Tài liệu thao tác theo bộ kit phiên bản 0.1.0; người bảo trì cập nhật ảnh và hướng dẫn khi giao diện thay đổi.')

# Page numbering supports a multi-page procedural manual.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Trang ');run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
OUT.parent.mkdir(parents=True,exist_ok=True)
# Remove Word template paragraph borders, including the built-in Title blue rule.
for tree in [doc.styles.element, doc._element]:
    for border in tree.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
